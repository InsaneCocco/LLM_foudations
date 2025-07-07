from docx import Document
from docx.shared import Inches
import os

# Create a new document
document = Document()

# Add a heading
document.add_heading('This is the Document Title', level=1)

# Add a paragraph with different text formatting
p = document.add_paragraph('This is a simple paragraph. ')
p.add_run('This part is bold.').bold = True
p.add_run(' And this part is italic.').italic = True

life_hack = '''To quickly chill a drink (or anything else small) in a hurry, wrap a wet paper towel around the can or bottle before placing it in the freezer.

The wet paper towel helps transfer the cold from the freezer to the drink much more efficiently, chilling it 
significantly faster than just putting the bare can/bottle in. Just remember to set a timer so you don't forget it and 
end up with an exploded can!'''
p.add_run(life_hack).bold = True

# Add a table
table_data = [
    ('ID', 'Product', 'Price'),
    (1, 'Keyboard', '$50'),
    (2, 'Mouse', '$25'),
    (3, 'Monitor', '$200')]

table = document.add_table(rows=1, cols=3)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'ID'
hdr_cells[1].text = 'Product'
hdr_cells[2].text = 'Price'
for item_id, product, price in table_data[1:]:
    row_cells = table.add_row().cells
    row_cells[0].text = str(item_id)
    row_cells[1].text = product
    row_cells[2].text = price

# Save the document
document.save(f'{os.getcwd()}/mydocument.docx')

